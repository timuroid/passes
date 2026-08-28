import io
import math
from datetime import date, datetime, time, timedelta, timezone
from typing import Literal
from zoneinfo import ZoneInfo

from fastapi import Depends, FastAPI, HTTPException, Query, Request, Response, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, func, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.config import Settings, get_settings
from app.database import build_engine, build_session_factory
from app.events import EventBroker
from app.models import AppSetting, PassRequest, SessionToken, User, utc_now
from app.schemas import (
    LoginRequest,
    PageMeta,
    PassPage,
    PassView,
    PublicPassCreate,
    UserView,
    UserCreate,
    UserActivationUpdate,
    UserPasswordUpdate,
    VisibilityUpdate,
    DriverThemeUpdate,
    PublicSettingsView,
)
from app.security import (
    AuthContext,
    SESSION_COOKIE,
    SlidingWindowLimiter,
    get_current_user,
    hash_session_token,
    new_csrf_token,
    new_session_token,
    request_client_key,
    require_admin,
    require_admin_csrf,
    require_csrf,
    verify_password,
    hash_password,
)
from app.vehicle_numbers import normalize_vehicle_search


def create_app(settings: Settings | None = None) -> FastAPI:
    settings = settings or get_settings()
    engine = build_engine(settings.database_url)
    session_factory = build_session_factory(engine)

    application = FastAPI(
        title="Транспортные средства ЗТЗ",
        version="1.0.0",
        docs_url="/api/docs",
        openapi_url="/api/openapi.json",
        redoc_url=None,
    )
    application.state.settings = settings
    application.state.engine = engine
    application.state.session_factory = session_factory
    application.state.events = EventBroker()
    application.state.login_limiter = SlidingWindowLimiter(limit=10, window_seconds=60)
    application.state.submit_limiter = SlidingWindowLimiter(limit=30, window_seconds=60)

    def get_db(request: Request):
        with request.app.state.session_factory() as db:
            yield db

    def build_pass_filters(
        date_from: date | None,
        date_to: date | None,
        visibility: Literal["visible", "hidden", "all"],
        search: str | None,
    ) -> list:
        filters = []
        local_tz = ZoneInfo(settings.app_timezone)
        if date_from:
            start_local = datetime.combine(date_from, time.min, tzinfo=local_tz)
            filters.append(PassRequest.submitted_at >= start_local.astimezone(timezone.utc))
        if date_to:
            end_local = datetime.combine(date_to + timedelta(days=1), time.min, tzinfo=local_tz)
            filters.append(PassRequest.submitted_at < end_local.astimezone(timezone.utc))
        if date_from and date_to and date_from > date_to:
            raise HTTPException(status_code=422, detail="Начальная дата не может быть позже конечной")
        if search:
            search_key = normalize_vehicle_search(search)
            if not search_key:
                raise HTTPException(status_code=422, detail="Введите буквы или цифры для поиска")
            filters.append(PassRequest.vehicle_number_search.contains(search_key, autoescape=True))

        if visibility == "visible":
            filters.append(PassRequest.is_hidden.is_(False))
        elif visibility == "hidden":
            filters.append(PassRequest.is_hidden.is_(True))
        return filters

    @application.middleware("http")
    async def same_origin_and_no_store(request: Request, call_next):
        if request.method not in {"GET", "HEAD", "OPTIONS"} and request.url.path.startswith("/api/"):
            origin = request.headers.get("origin")
            if origin:
                forwarded_proto = request.headers.get("x-forwarded-proto", request.url.scheme)
                expected = f"{forwarded_proto}://{request.headers.get('host', request.url.netloc)}"
                if origin.rstrip("/") != expected.rstrip("/"):
                    return Response("Недопустимый источник запроса", status_code=403)
        response = await call_next(request)
        if request.url.path.startswith("/api/"):
            response.headers.setdefault("Cache-Control", "no-store")
        return response

    @application.get("/api/health", tags=["system"])
    def health(db: Session = Depends(get_db)) -> dict[str, str]:
        db.execute(text("SELECT 1"))
        return {"status": "ok"}

    @application.post("/api/auth/login", tags=["auth"])
    async def login(payload: LoginRequest, request: Request, response: Response, db: Session = Depends(get_db)):
        username = payload.username.strip().lower()
        limiter_key = f"{request_client_key(request)}:{username}"
        if not await request.app.state.login_limiter.allow(limiter_key):
            raise HTTPException(
                status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                detail="Слишком много попыток входа. Повторите через минуту.",
                headers={"Retry-After": "60"},
            )

        user = db.scalar(select(User).where(User.username == username))
        if user is None or not user.is_active or not verify_password(user.password_hash, payload.password):
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Неверный логин или пароль")

        now = utc_now()
        db.execute(delete(SessionToken).where(SessionToken.expires_at <= now))
        raw_token = new_session_token()
        csrf_token = new_csrf_token()
        expires_at = now + timedelta(hours=settings.session_ttl_hours)
        db.add(
            SessionToken(
                token_hash=hash_session_token(raw_token),
                csrf_token=csrf_token,
                user_id=user.id,
                expires_at=expires_at,
            )
        )
        db.commit()
        response.set_cookie(
            key=SESSION_COOKIE,
            value=raw_token,
            max_age=settings.session_ttl_hours * 3600,
            httponly=True,
            secure=settings.cookie_secure,
            samesite="strict",
            path="/",
        )
        return {
            "user": UserView.model_validate(user),
            "csrf_token": csrf_token,
            "expires_at": expires_at,
        }

    @application.get("/api/auth/me", tags=["auth"])
    def me(current: AuthContext = Depends(get_current_user)) -> dict[str, object]:
        return {
            "user": {
                "id": current.user_id,
                "username": current.username,
                "role": current.role,
            },
            "csrf_token": current.csrf_token,
        }

    @application.post("/api/auth/logout", status_code=204, tags=["auth"])
    def logout(
        request: Request,
        response: Response,
        _: AuthContext = Depends(require_csrf),
        db: Session = Depends(get_db),
    ) -> Response:
        raw_token = request.cookies.get(SESSION_COOKIE)
        if raw_token:
            db.execute(delete(SessionToken).where(SessionToken.token_hash == hash_session_token(raw_token)))
            db.commit()
        response.delete_cookie(SESSION_COOKIE, path="/", samesite="strict", secure=settings.cookie_secure)
        response.status_code = 204
        return response

    @application.post("/api/public/passes", response_model=PassView, status_code=201, tags=["public"])
    async def create_pass(payload: PublicPassCreate, request: Request, db: Session = Depends(get_db)):
        if not await request.app.state.submit_limiter.allow(request_client_key(request)):
            raise HTTPException(
                status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                detail="Слишком много отправок. Повторите через минуту.",
                headers={"Retry-After": "60"},
            )
        item = PassRequest(
            vehicle_number=payload.vehicle_number,
            vehicle_number_search=normalize_vehicle_search(payload.vehicle_number),
            phone_number=payload.phone_number,
        )
        db.add(item)
        db.commit()
        db.refresh(item)
        await request.app.state.events.publish({"type": "pass.created", "id": item.id})
        return item

    @application.get("/api/public/settings", response_model=PublicSettingsView, tags=["public"])
    def public_settings(db: Session = Depends(get_db)) -> PublicSettingsView:
        setting = db.get(AppSetting, "driver_theme")
        theme = setting.value if setting and setting.value in {"light", "dark"} else "light"
        return PublicSettingsView(driver_theme=theme)

    @application.patch("/api/settings/driver-theme", response_model=PublicSettingsView, tags=["admin"])
    async def update_driver_theme(
        payload: DriverThemeUpdate,
        request: Request,
        _: AuthContext = Depends(require_admin_csrf),
        db: Session = Depends(get_db),
    ) -> PublicSettingsView:
        setting = db.get(AppSetting, "driver_theme")
        if setting is None:
            setting = AppSetting(key="driver_theme", value=payload.theme)
            db.add(setting)
        else:
            setting.value = payload.theme
        db.commit()
        await request.app.state.events.publish({"type": "settings.driver_theme", "theme": payload.theme})
        return PublicSettingsView(driver_theme=payload.theme)

    @application.get("/api/passes", response_model=PassPage, tags=["passes"])
    def list_passes(
        date_from: date | None = None,
        date_to: date | None = None,
        sort: Literal["asc", "desc"] = "desc",
        page: int = Query(1, ge=1),
        page_size: int = Query(25, ge=10, le=100),
        visibility: Literal["visible", "hidden", "all"] = "visible",
        search: str | None = Query(None, min_length=1, max_length=24),
        _: AuthContext = Depends(get_current_user),
        db: Session = Depends(get_db),
    ) -> PassPage:
        filters = build_pass_filters(date_from, date_to, visibility, search)

        total = db.scalar(select(func.count()).select_from(PassRequest).where(*filters)) or 0
        ordering = PassRequest.submitted_at.asc() if sort == "asc" else PassRequest.submitted_at.desc()
        id_ordering = PassRequest.id.asc() if sort == "asc" else PassRequest.id.desc()
        items = list(
            db.scalars(
                select(PassRequest)
                .where(*filters)
                .order_by(ordering, id_ordering)
                .offset((page - 1) * page_size)
                .limit(page_size)
            )
        )
        pages = math.ceil(total / page_size) if total else 0
        return PassPage(
            items=[PassView.model_validate(item) for item in items],
            meta=PageMeta(page=page, page_size=page_size, total=total, pages=pages, sort=sort),
        )

    @application.get("/api/passes/export.xlsx", tags=["passes"])
    def export_passes(
        date_from: date | None = None,
        date_to: date | None = None,
        sort: Literal["asc", "desc"] = "desc",
        visibility: Literal["visible", "hidden", "all"] = "visible",
        search: str | None = Query(None, min_length=1, max_length=24),
        _: AuthContext = Depends(get_current_user),
        db: Session = Depends(get_db),
    ) -> Response:
        filters = build_pass_filters(date_from, date_to, visibility, search)
        ordering = PassRequest.submitted_at.asc() if sort == "asc" else PassRequest.submitted_at.desc()
        id_ordering = PassRequest.id.asc() if sort == "asc" else PassRequest.id.desc()
        items = list(db.scalars(select(PassRequest).where(*filters).order_by(ordering, id_ordering)))

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Транспортные средства"
        headers = ["ID", "Номер автомобиля", "Телефон", "Получен", "Статус"]
        sheet.append(headers)
        header_fill = PatternFill("solid", fgColor="FF3C00")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        local_tz = ZoneInfo(settings.app_timezone)
        for item in items:
            submitted_at = item.submitted_at
            if submitted_at.tzinfo is None:
                submitted_at = submitted_at.replace(tzinfo=timezone.utc)
            sheet.append(
                [
                    item.id,
                    item.vehicle_number,
                    item.phone_number or "",
                    submitted_at.astimezone(local_tz).replace(tzinfo=None),
                    "Скрыт" if item.is_hidden else "Активен",
                ]
            )

        sheet.column_dimensions["A"].width = 10
        sheet.column_dimensions["B"].width = 24
        sheet.column_dimensions["C"].width = 22
        sheet.column_dimensions["D"].width = 22
        sheet.column_dimensions["E"].width = 14
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            row[1].number_format = "@"
            row[2].number_format = "@"
            row[3].number_format = "DD.MM.YYYY HH:MM"

        output = io.BytesIO()
        workbook.save(output)
        filename = f"vehicles-{datetime.now(local_tz):%Y-%m-%d}.xlsx"
        return Response(
            output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    @application.patch("/api/passes/{pass_id}/visibility", response_model=PassView, tags=["passes"])
    async def update_visibility(
        pass_id: int,
        payload: VisibilityUpdate,
        request: Request,
        current: AuthContext = Depends(require_csrf),
        db: Session = Depends(get_db),
    ) -> PassRequest:
        item = db.get(PassRequest, pass_id)
        if item is None:
            raise HTTPException(status_code=404, detail="Запись не найдена")
        if payload.hidden and not item.is_hidden:
            item.is_hidden = True
            item.hidden_at = utc_now()
            item.hidden_by_user_id = current.user_id
            item.restored_at = None
        elif not payload.hidden and item.is_hidden:
            item.is_hidden = False
            item.restored_at = utc_now()
            item.hidden_at = None
            item.hidden_by_user_id = None
        db.commit()
        db.refresh(item)
        await request.app.state.events.publish(
            {"type": "pass.visibility", "id": item.id, "hidden": item.is_hidden}
        )
        return item

    @application.get("/api/users", response_model=list[UserView], tags=["admin"])
    def list_users(
        _: AuthContext = Depends(require_admin),
        db: Session = Depends(get_db),
    ) -> list[User]:
        return list(db.scalars(select(User).order_by(User.username.asc())))

    @application.post("/api/users", response_model=UserView, status_code=201, tags=["admin"])
    def create_user(
        payload: UserCreate,
        _: AuthContext = Depends(require_admin_csrf),
        db: Session = Depends(get_db),
    ) -> User:
        user = User(
            username=payload.username,
            password_hash=hash_password(payload.password),
            role=payload.role,
            is_active=True,
        )
        db.add(user)
        try:
            db.commit()
        except IntegrityError as error:
            db.rollback()
            raise HTTPException(status_code=409, detail="Пользователь с таким логином уже существует") from error
        db.refresh(user)
        return user

    @application.patch("/api/users/{user_id}/active", response_model=UserView, tags=["admin"])
    async def update_user_active(
        user_id: int,
        payload: UserActivationUpdate,
        request: Request,
        current: AuthContext = Depends(require_admin_csrf),
        db: Session = Depends(get_db),
    ) -> User:
        user = db.get(User, user_id)
        if user is None:
            raise HTTPException(status_code=404, detail="Пользователь не найден")

        if not payload.active and user.is_active:
            if user.id == current.user_id:
                raise HTTPException(status_code=409, detail="Нельзя отключить собственную активную учётную запись")
            if user.role == "admin":
                active_admins = db.scalar(
                    select(func.count()).select_from(User).where(User.role == "admin", User.is_active.is_(True))
                ) or 0
                if active_admins <= 1:
                    raise HTTPException(status_code=409, detail="Нельзя отключить последнего активного администратора")
            user.is_active = False
            db.execute(delete(SessionToken).where(SessionToken.user_id == user.id))
        elif payload.active and not user.is_active:
            user.is_active = True

        db.commit()
        db.refresh(user)
        await request.app.state.events.publish({"type": "user.active", "id": user.id, "active": user.is_active})
        return user

    @application.patch("/api/users/{user_id}/password", response_model=UserView, tags=["admin"])
    async def update_user_password(
        user_id: int,
        payload: UserPasswordUpdate,
        request: Request,
        _: AuthContext = Depends(require_admin_csrf),
        db: Session = Depends(get_db),
    ) -> User:
        user = db.get(User, user_id)
        if user is None:
            raise HTTPException(status_code=404, detail="Пользователь не найден")
        user.password_hash = hash_password(payload.password)
        db.execute(delete(SessionToken).where(SessionToken.user_id == user.id))
        db.commit()
        db.refresh(user)
        await request.app.state.events.publish({"type": "user.password", "id": user.id})
        return user

    @application.get("/api/events", tags=["events"])
    async def events(request: Request, _: AuthContext = Depends(get_current_user)) -> StreamingResponse:
        queue = await request.app.state.events.subscribe()
        raw_token = request.cookies.get(SESSION_COOKIE, "")
        token_hash = hash_session_token(raw_token)

        def session_is_active() -> bool:
            with request.app.state.session_factory() as db:
                session = db.scalar(select(SessionToken).where(SessionToken.token_hash == token_hash))
                if session is None or not session.user.is_active:
                    return False
                expires_at = session.expires_at
                if expires_at.tzinfo is None:
                    expires_at = expires_at.replace(tzinfo=timezone.utc)
                return expires_at > datetime.now(timezone.utc)

        return StreamingResponse(
            request.app.state.events.stream(queue, session_is_active),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache, no-transform",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no",
            },
        )

    return application


app = create_app()
